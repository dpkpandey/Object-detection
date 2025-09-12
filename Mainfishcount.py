#!/usr/bin/env python3
import cv2
import threading
import datetime
import time
import numpy as np
import os
import queue
from ultralytics import YOLO
from sort import Sort
from openpyxl import Workbook, load_workbook
import cvzone

# =================== Config ===================
DROP_INTERVAL = float('inf') #90
QUEUE_MAXSIZE = 3000
EXCEL_UPDATE_INTERVAL = 600  # seconds
DISPLAY_WINDOW = "High-FPS Fish Counter (Video)"
MODEL_PATH = "last1.engine"
VIDEO_PATH = "rt.avi"
DEFAULT_LINE = [440, 300, 920, 300]
CLASS_LIST = ['fish']

start_time = time.time()

class SingleCamFishCounter:
    def __init__(self, model_path, video_source, line=None, footage_recorder=None):
        self.model = YOLO(model_path)
        self.video_source = video_source
        self.frame_queue = queue.Queue(maxsize=QUEUE_MAXSIZE)
        self.tracker = Sort(max_age=30, min_hits=3, iou_threshold=0.3)
        self.counter = set()
        self.capture_counter = 0
        self.object_count = 0
        self.lock = threading.Lock()
        self.frame = None
        self.prev_gray = None
        self.stopped = False
        self.paused = False  # NEW: pause flag
        self.last_update_time = time.time()
        self.last_log_time = time.time()
        self.processed_frames = 0

        # Counting line: editable at runtime
        self.line = line if line is not None else list(DEFAULT_LINE)

        # External footage recorder (if provided by GUI)
        self.footage_recorder = footage_recorder

        # Backward compatibility: internal writers only if no external recorder
        if self.footage_recorder is None:
            timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
            self.writer = cv2.VideoWriter(
                f"Processed_SingleCam_{timestamp}.avi",
                cv2.VideoWriter_fourcc(*'XVID'), 90, (640, 480))
            self.raw_writer = cv2.VideoWriter(
                f"Raw_SingleCam_{timestamp}.avi",
                cv2.VideoWriter_fourcc(*'XVID'), 110, (1280, 720))
        else:
            self.writer = None
            self.raw_writer = None

        # ======== NEW: GUI-controllable imshow toggle ========
        self.imshow_enabled = True   # default ON (existing behavior)
        self._window_open = False    # track window state

    # ==== Pause and Resume ====
    def pause(self):
        self.paused = True

    def resume(self):
        self.paused = False

    # ==== NEW: Allow GUI to toggle preview window ====
    def set_imshow(self, enabled: bool):
        """Enable/disable cv2.imshow from the GUI at runtime."""
        self.imshow_enabled = bool(enabled)

    def show_window(self):
        """Show the preview window."""
        self.set_imshow(True)

    def hide_window(self):
        """Hide the preview window without stopping the program."""
        self.set_imshow(False)

    def frame_is_bright(self, frame, min_mean=10, min_std=3):
        gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
        return np.mean(gray) > min_mean and np.std(gray) > min_std

    def capture_thread(self):
        cap = cv2.VideoCapture(self.video_source)
        while not self.stopped:
            # PAUSE SUPPORT
            if self.paused:
                time.sleep(0.05)
                continue

            ret, frame = cap.read()
            if not ret:
                print("[INFO] End of video file.")
                self.stopped = True
                break

            # Write raw if internal writer exists
            if self.raw_writer is not None:
                self.raw_writer.write(frame)

            # External recorder (raw)
            if self.footage_recorder is not None:
                self.footage_recorder.record_frame(raw_frame=frame)

            self.capture_counter += 1
            if self.capture_counter % DROP_INTERVAL == 0:
                continue

            if not self.frame_is_bright(frame):
                continue

            #if self.frame_queue.qsize() > 2000:
             #   continue  # skip if overloaded

            try:
                self.frame_queue.put_nowait(frame)
            except queue.Full:
                pass
        cap.release()

    def process_thread(self):
        processed_frames = 0
        while not self.stopped:
            # PAUSE SUPPORT
            if self.paused:
                time.sleep(0.05)
                continue

            try:
                frame = self.frame_queue.get(timeout=1)
            except queue.Empty:
                continue

            try:
                detections, tracks, crossing = self.process_frame(frame)
            except Exception as e:
                print(f"[ERROR] Inference failed: {e}")
                continue

            self.object_count = len(self.counter)
            color = (0, 255, 0) if crossing else (0, 0, 255)
            cv2.line(frame, tuple(self.line[:2]), tuple(self.line[2:]), color, 10)

            for result in tracks:
                x1, y1, x2, y2, track_id = map(int, result)
                cx, cy = x1 + (x2 - x1) // 2, y1 + (y2 - y1) // 2
                cv2.rectangle(frame, (x1, y1), (x2, y2), (0, 255, 255), 2)
                cv2.circle(frame, (cx, cy), 10, (127, 0, 0), -1)
                cvzone.putTextRect(frame, f'{track_id}', [x1 + 8, y1 - 12],
                                   colorR=(0, 0, 255), thickness=2, scale=1.5)

            cvzone.putTextRect(frame, f'Detected fish No = {self.object_count}', [80, 34],
                               colorR=(0, 0, 255), thickness=4, scale=2.3, border=3)
            cvzone.putTextRect(frame, f'MY FISH COUNTER', [1040, 700],
                               colorR=(255, 0, 100), thickness=2, scale=1.3, border=1)

            resized = cv2.resize(frame, (640, 480))
            with self.lock:
                self.frame = resized

            # Internal writer (processed)
            if self.writer is not None:
                self.writer.write(resized)

            # External recorder (processed)
            if self.footage_recorder is not None:
                self.footage_recorder.record_frame(processed_frame=resized)

            processed_frames += 1
            self.processed_frames += 1

            if time.time() - self.last_log_time > 5:
                total_elapsed = time.time() - start_time
                fps = processed_frames / total_elapsed if total_elapsed > 0 else 0
                print(f"[INFO] Fish count: {self.object_count}, Frames: {processed_frames}, FPS: {fps:.2f}")
                self.last_log_time = time.time()

        # Release internal writers if they exist
        if self.writer is not None:
            self.writer.release()
        if self.raw_writer is not None:
            self.raw_writer.release()

    def process_frame(self, frame):
        results = self.model(frame, device=0)
        detections = np.empty((0, 5))

        for info in results:
            for box in info.boxes:
                x1, y1, x2, y2 = box.xyxy[0].cpu().numpy()
                conf = box.conf[0].cpu().item()
                class_id = int(box.cls[0].cpu().item())
                if CLASS_LIST[class_id] == 'fish' and conf > 0.25:
                    detections = np.vstack((detections, [x1, y1, x2, y2, conf]))

        tracks = self.tracker.update(detections)
        crossing = set()
        for result in tracks:
            x1, y1, x2, y2, track_id = map(int, result)
            cx, cy = x1 + (x2 - x1) // 2, y1 + (y2 - y1) // 2
            if self.line[0] < cx < self.line[2] and (cy - 90) < self.line[1] < (cy + 25):
                if track_id not in self.counter:
                    self.counter.add(track_id)
                    crossing.add(track_id)
        return detections, tracks, crossing

    def update_excel(self):
        filename = 'fish_count_log_video.xlsx'
        backup = filename.replace('.xlsx', '_backup.xlsx')
        timestamp = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')

        try:
            if os.path.exists(filename):
                wb = load_workbook(filename)
                ws = wb.active
            else:
                wb = Workbook()
                ws = wb.active
                ws.append(["Timestamp", "Fish Count"])
            ws.append([timestamp, len(self.counter)])
            wb.save(filename)
        except PermissionError:
            try:
                wb.save(backup)
                print(f"[INFO] Excel backup saved at {timestamp}")
            except Exception as e:
                print(f"[ERROR] Backup Excel save failed: {e}")

    def display_thread(self):
        while not self.stopped:
            # PAUSE SUPPORT
            if self.paused:
                time.sleep(0.05)
                continue

            # ======== MODIFIED: respect GUI-controlled imshow toggle ========
            if not self.imshow_enabled:
                # if a window is open, close it once; then idle
                if self._window_open:
                    try:
                        cv2.destroyWindow(DISPLAY_WINDOW)
                    except Exception:
                        pass
                    self._window_open = False
                time.sleep(0.05)
                continue
            # =================================================================

            with self.lock:
                frame = self.frame.copy() if self.frame is not None else None

            if frame is None:
                frame = np.ones((480, 640, 3), dtype=np.uint8) * 255
                cv2.putText(frame, "Waiting...", (100, 240), cv2.FONT_HERSHEY_SIMPLEX, 1, (0, 0, 255), 2)

            cv2.imshow(DISPLAY_WINDOW, frame)  # FIXED: Uncommented this line
            self._window_open = True

            if time.time() - self.last_update_time >= EXCEL_UPDATE_INTERVAL:
                self.update_excel()
                self.last_update_time = time.time()

            # ======== MODIFIED: Handle window close and ESC key ========
            key = cv2.waitKey(1) & 0xFF
            
            # Check if window was closed by user (X button)
            try:
                if cv2.getWindowProperty(DISPLAY_WINDOW, cv2.WND_PROP_VISIBLE) < 1:
                    self.imshow_enabled = False  # Disable imshow but don't stop program
                    continue
            except:
                # Window doesn't exist, disable imshow
                self.imshow_enabled = False
                continue
                
            # ESC key still terminates the program
            if key == 27:  # ESC key
                self.stopped = True
                break
            # ============================================================

        # cleanup (also closes if window was disabled earlier)
        try:
            cv2.destroyAllWindows()
        except Exception:
            pass

    def run(self):
        threads = [
            threading.Thread(target=self.capture_thread, daemon=True),
            threading.Thread(target=self.process_thread, daemon=True),
            threading.Thread(target=self.display_thread, daemon=True)
        ]
        for t in threads:
            t.start()

        while not self.stopped:
            time.sleep(0.1)

        total_time = time.time() - start_time
        fps = self.processed_frames / total_time if total_time > 0 else 0

        print("\n=== Final Fish Count Summary ===")
        print(f"Total detected fish: {len(self.counter)}")
        print(f"Total time: {time.time() - start_time:.2f} seconds")
        print(f"[INFO] Fish Count : {self.object_count}, Frames: {self.processed_frames}, FPS :{fps:.2f}")

# Main (kept for standalone tests; your GUI can import the class and call set_imshow)
if __name__ == "__main__":
    counter = SingleCamFishCounter(MODEL_PATH, VIDEO_PATH)
    # Example: disable on-screen preview (GUI will normally call this)
    # counter.set_imshow(False)
    counter.run()
